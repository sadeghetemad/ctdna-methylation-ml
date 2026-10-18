"""
10_report.py: Word report of the step 08 and 09 results, for technical and non-technical readers.

Reads results.json, the tables and the figures written by 08_ml_analysis.py and writes
ML_Report.docx next to them.

Example:
    python scripts/10_report.py --results results/ml/nfcore_paper --data notebooks/data/nfcore_full
"""

import argparse
import json
import time
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
NAVY, TEAL, MUTED = RGBColor(0x0A, 0x2F, 0x5C), RGBColor(0x00, 0x8A, 0x84), RGBColor(0x6B, 0x72, 0x80)


def pct(x):
    return f"{100 * x:.0f}%"


def ci(lo, hi, f=pct):
    return f"{f(lo)}–{f(hi)}"


def num(x, d=2):
    return f"{x:.{d}f}"


class Report:
    def __init__(self):
        self.doc = Document()
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = "Calibri", Pt(10.5)
        st.paragraph_format.space_after = Pt(5)
        for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
            h = self.doc.styles[name]
            h.font.name, h.font.size, h.font.color.rgb, h.font.bold = "Calibri", Pt(size), NAVY, True

    def h(self, text, level=1):
        self.doc.add_heading(text, level=level)

    def p(self, text, bold=False, italic=False, color=None, size=None):
        par = self.doc.add_paragraph()
        r = par.add_run(text)
        r.bold, r.italic = bold, italic
        if color is not None:
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)
        return par

    def bullets(self, items):
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def box(self, title, lines, fill="E6F6F5"):
        t = self.doc.add_table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shd)
        cell.paragraphs[0].add_run(title).bold = True
        for ln in lines:
            cell.add_paragraph(ln, style="List Bullet")
        self.doc.add_paragraph()

    def table(self, df, widths=None, fontsize=8.5):
        t = self.doc.add_table(rows=1, cols=len(df.columns))
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, name in zip(t.rows[0].cells, df.columns):
            c.text = str(name)
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for c, v in zip(cells, row):
                c.text = str(v)
        for row in t.rows:
            for i, c in enumerate(row.cells):
                for par in c.paragraphs:
                    for r in par.runs:
                        r.font.size = Pt(fontsize)
                if widths:
                    c.width = Cm(widths[i])
        self.doc.add_paragraph()

    def figure(self, path, caption, width=16):
        if Path(path).exists():
            self.doc.add_picture(str(path), width=Cm(width))
            self.p(caption, italic=True, color=MUTED, size=9)


def batch_figure(path, sel, ens, batches, scorable, test_cancers, cov_only, cv_within):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    rows = [("Test set, all", sel["auc_test"], ens["auc_test"], "")]
    for b in batches:
        ok = b in scorable
        rows.append((f"Test set, batch {b}", sel[f"auc_test_{b}"] if ok else None, ens[f"auc_test_{b}"] if ok else None,
                     "" if ok else f"n.e. ({int(test_cancers.get(b, 0))} cancer)"))
    for b, v in cv_within.items():
        rows.append((f"CV inside batch {b} only", v, None, ""))
    fig, ax = plt.subplots(figsize=(7.2, 0.45 * len(rows) + 1.0))
    yy = list(range(len(rows)))[::-1]
    for y_, (lab, a, e, note) in zip(yy, rows):
        if a is not None:
            ax.barh(y_ + 0.18, a - 0.5, left=0.5, height=0.34, color="#eb6834")
            ax.text(max(a, 0.5) + 0.005, y_ + 0.18, f"{a:.2f}", va="center", fontsize=7.5)
        if e is not None:
            ax.barh(y_ - 0.18, e - 0.5, left=0.5, height=0.34, color="#8a5cc2")
            ax.text(max(e, 0.5) + 0.005, y_ - 0.18, f"{e:.2f}", va="center", fontsize=7.5)
        if note:
            ax.text(0.505, y_, note, va="center", fontsize=7.5, color="#6b7280")
    ax.set_yticks(yy, [r[0] for r in rows], fontsize=8)
    ax.axvline(0.5, color="black", linewidth=1)
    ax.axvline(cov_only, color="#8a8985", linestyle="--", linewidth=1)
    ax.set_xlim(min(0.45, min(v for r in rows for v in r[1:3] if v is not None) - 0.05), 1.0)
    ax.set_xlabel("AUC (orange: selected model; purple: ensemble; 0.5 = chance)")
    ax.set_title("Performance overall and per sequencing batch", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=130)
    plt.close(fig)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default="results/ml/nfcore_paper")
    ap.add_argument("--data", default="notebooks/data/nfcore_full", help="Step 07 folder (for per-batch counts).")
    args = ap.parse_args()
    res = Path(args.results) if Path(args.results).is_absolute() else ROOT / args.results
    R = json.loads((res / "results.json").read_text())
    d, sel, ens = R["data"], R["heldout"]["selected"], R["heldout"]["ensemble"]
    perm, s = R["permutation"], R["selection"]
    nested = pd.DataFrame(R["nested"])
    ho = pd.DataFrame(R["heldout_all"])
    conf = pd.DataFrame(R["confounders"]).T
    uni = pd.DataFrame(R["univariate"]).T
    rob = pd.DataFrame(R["robustness"]).T
    batches = sorted(d["batches"])
    big = d["big_batch"]
    sens100 = round(100 * sel["sens_at_95spec"])
    fp100 = round(100 * (1 - sel["spec_reached"]))
    cov_only = rob.iloc[0]["auc_test"]
    tech_max = conf.drop(index="Mean region methylation", errors="ignore")["abs_auc"].astype(float).max()
    bench = json.loads((res / "benchmark.json").read_text()) if (res / "benchmark.json").exists() else None
    bench_line = None
    if bench:
        a = bench["all"]
        ours = a["our_sens_at_published_spec"]
        bench_line = (f"Benchmark: on the same test samples the published ELSA-seq classifier detects {pct(a['published_sens'])} "
                      f"of Stage I cancers at {pct(a['published_spec'])} specificity; at that specificity our model detects "
                      f"{pct(ours)}.")
        weak = [b for b in batches if b in bench and bench[b]["published_sens"] < a["published_sens"] - 0.1]
        if weak:
            bench_line += (" The published classifier is weaker in batch "
                           + ", ".join(f"{b} ({pct(bench[b]['published_sens'])} sensitivity)" for b in weak) + ".")
    data_dir = Path(args.data) if Path(args.data).is_absolute() else ROOT / args.data
    smp = pd.read_csv(data_dir / "samples.tsv", sep="\t")
    DEV, TEST = d.get("train_group", "TV"), d.get("test_group", "SBT")
    tes = smp[smp["model_group"] == TEST]
    devs = smp[smp["model_group"] == DEV]
    small = min(batches, key=lambda b: d["batches"][b])
    share_dev = (devs["submission_series"] == small).mean()
    share_test = (tes["submission_series"] == small).mean()
    gap = sel["auc_test"] - s["best_nested_auc"]
    if gap > 0.05:
        comp_line = (f"The test result (AUC {num(sel['auc_test'])}) is higher than the development estimate "
                     f"(nested AUC {num(s['best_nested_auc'])}). The small, harder batch {small} makes up "
                     f"{share_dev:.0%} of the development set but {share_test:.0%} of the test set, so the test numbers "
                     "may be optimistic for a mixed population; the nested estimate is the conservative one.")
    elif gap < -0.05:
        comp_line = (f"The test result (AUC {num(sel['auc_test'])}) is lower than the development estimate "
                     f"(nested AUC {num(s['best_nested_auc'])}). The small, harder batch {small} makes up "
                     f"{share_dev:.0%} of the development set but {share_test:.0%} of the test set, so the test set is "
                     "the harder, more conservative check.")
    else:
        comp_line = (f"The test result (AUC {num(sel['auc_test'])}) agrees with the development estimate "
                     f"(nested AUC {num(s['best_nested_auc'])}), which suggests the model generalises within this study.")
    test_cancers = tes.groupby("submission_series")["ml_label"].sum().to_dict()
    benign = smp[smp["stage"] == "Benign"].groupby("submission_series").size().to_dict()
    MIN_CANCERS = 5
    scorable = [b for b in batches if test_cancers.get(b, 0) >= MIN_CANCERS]
    unscorable = [b for b in batches if b not in scorable]
    cv_within = {b: rob.loc[[i for i in rob.index if f"batch {b} only" in i][0], "cv_auc"]
                 for b in batches if any(f"batch {b} only" in i for i in rob.index)}
    dev_cancers = devs.groupby("submission_series")["ml_label"].sum().to_dict()
    works = [b for b in scorable if sel[f"auc_test_{b}"] >= 0.75]
    fails = [b for b in scorable if sel[f"auc_test_{b}"] < 0.65]
    parts = []
    if works:
        parts.append("Within batch " + ", ".join(f"{b} ({int((tes['submission_series'] == b).sum())} test samples)" for b in works)
                     + " the selected model reaches AUC " + ", ".join(num(sel[f"auc_test_{b}"]) for b in works)
                     + " (ensemble " + ", ".join(num(ens[f"auc_test_{b}"]) for b in works) + ").")
    for b in fails:
        parts.append(f"In batch {b} it does not work (AUC {num(sel[f'auc_test_{b}'])}), although cross-validation inside "
                     f"{b} alone reaches AUC {num(cv_within.get(b, float('nan')))}: the signal exists in that batch, but a model "
                     f"trained on the development set, which holds only {int(dev_cancers.get(b, 0))} cancer(s) from {b}, does "
                     "not transfer to it.")
    for b in unscorable:
        parts.append(f"Batch {b} has only {int(test_cancers.get(b, 0))} cancer(s) in the test set and cannot be scored there; "
                     f"cross-validation inside it alone gives AUC {num(cv_within.get(b, float('nan')))}.")
    batch_line = " ".join(parts)

    rp = Report()
    rp.p("Early Lung Cancer Detection from ctDNA Methylation", bold=True, size=20, color=NAVY)
    rp.p(f"Machine-learning results report  ·  {time.strftime('%d %B %Y')}", color=MUTED, size=11)

    # ------------------------------------------------------------ executive summary
    rp.h("Executive summary")
    rp.p(f"We trained and tested models that tell Stage I lung cancer apart from people without cancer, using DNA "
         f"methylation measured in blood plasma. The data are {d['samples']} people from a published study "
         f"({d['cancer']} with Stage I lung cancer and {d['control']} without cancer). Models were built on one part of the "
         f"data ({d['train']} people) and then tested once on a separate part ({d['test']} people) that was never used "
         f"while building them.")
    rp.box("Key results on the independent test set", [
        f"The selected model ({s['best_candidate']}, using {s['best_input'].lower()}) separates cancer from no cancer with "
        f"an AUC of {num(sel['auc_test'])} (95% CI {ci(sel['auc_ci_low'], sel['auc_ci_high'], num)}). An AUC of 0.5 is "
        f"a coin flip and 1.0 is perfect.",
        f"At the threshold fixed on the development data for 5% false alarms, it detects about {sens100} of every 100 "
        f"Stage I cancers (95% CI {ci(sel['sens_ci_low'], sel['sens_ci_high'])}); on the test set this threshold gave "
        f"{fp100}% false alarms.",
        "The pre-specified ensemble of the three best models "
        + ("does better" if ens["auc_test"] > sel["auc_test"] + 0.01 else "does worse" if ens["auc_test"] < sel["auc_test"] - 0.01
           else "performs about the same")
        + f": AUC {num(ens['auc_test'])} (95% CI {ci(ens['auc_ci_low'], ens['auc_ci_high'], num)}), about "
        f"{round(100 * ens['sens_at_95spec'])} of 100 Stage I cancers detected at its threshold "
        f"({round(100 * (1 - ens['spec_reached']))}% false alarms on the test set).",
        f"Permutation test: with shuffled labels the same configuration averages AUC {num(perm['null_mean'])}, against "
        f"{num(perm['real_auc'])} with the real labels (p = {perm['p']:.3f} with {perm['n']} shuffles"
        + (", the smallest attainable)." if perm["p"] <= 1 / (perm["n"] + 1) + 1e-12 else ")."),
        batch_line,
        f"Technical and demographic variables alone (batch, age, gender, DNA input, depth) reach only AUC {num(cov_only)}.",
        *([bench_line] if bench_line else []),
    ])
    rp.p("What this means", bold=True)
    rp.bullets([
        "Methylation in blood carries a clear signal of early lung cancer in this dataset.",
        comp_line,
        "This is a baseline, not a screening test: screening needs high sensitivity at very high specificity, and "
        "confirmation on data from other hospitals and laboratories.",
    ])

    # ------------------------------------------------------------ data
    rp.h("1. Data and study design")
    rp.table(pd.DataFrame([
        ["Samples", f"{d['samples']} ({d['cancer']} Stage I cancer, {d['control']} controls)"],
        ["Stage I subtypes", ", ".join(f"{k}: {v}" for k, v in d["stage"].items())],
        ["Features", f"{d['regions']} methylation regions (ELSA panel) + read-level fragment features"],
        [f"Development set ({DEV}, training/validation)", f"{d['train']} samples, {d['train_cancer']} cancers"],
        [f"Test set ({TEST}, single-blind test)", f"{d['test']} samples, {d['test_cancer']} cancers, used once"],
        ["Sequencing batches", ", ".join(f"{k}: {v}" for k, v in d["batches"].items())],
    ], columns=["Item", "Value"]), widths=[5, 11.5])
    bl = pd.DataFrame(d["batch_by_label"])
    rp.p(f"The share of cancers differs between the two sequencing batches "
         + ", ".join(f"{b}: {int(bl.loc[b, '1'] if '1' in bl.columns else bl.loc[b].iloc[-1])} cancers of "
                     f"{int(bl.loc[b].sum())}" for b in bl.index)
         + ". A model could therefore learn the batch instead of cancer; this is tested explicitly below.")

    # ------------------------------------------------------------ signal
    rp.h("2. Is there a real signal before any modelling?")
    m = uni.loc["methylation"]
    rp.p(f"Each region was first tested on its own, on all samples (descriptive only; the models do not use these "
         f"results). {int(m['fdr_sig'])} of {int(m['features'])} methylation regions "
         f"({m['pct_sig']:.0f}%) differ between cancer and control after correcting for multiple testing (FDR < 5%), and "
         f"{int(m['up_in_cancer'])} of them are higher in cancer, the expected direction for this cancer-specific panel. "
         f"The best single region reaches AUC {num(m['best_auc'])}, so a combination of regions is needed. The region AUCs "
         f"inside batch {big} alone match the overall ones (r = {num(m['batch_consistency'])}), so the signal is not explained by the batch.")
    rp.figure(res / "fig02_univariate.png", "Figure 1. Left: methylation difference and significance of every region "
              "(orange: FDR < 5%). Right: region AUC on all samples against the same AUC within one batch.")
    rp.table(uni.reset_index().rename(columns={"index": "feature set"}).assign(
        pct_sig=lambda x: x["pct_sig"].map(lambda v: f"{v:.0f}%"), best_auc=lambda x: x["best_auc"].map(num),
        batch_consistency=lambda x: x["batch_consistency"].map(num))[
        ["feature set", "features", "fdr_sig", "pct_sig", "up_in_cancer", "down_in_cancer", "best_auc", "batch_consistency"]],
        fontsize=8)
    rp.figure(res / "fig03_pca.png", "Figure 2. Principal components of region methylation, coloured by label (left) and "
              "by sequencing batch (right).")
    rp.figure(res / "fig01_confounders.png", f"Figure 3. Each variable used alone as a predictor. Technical and demographic "
              f"variables reach at most AUC {num(tech_max)}; mean methylation (orange) is the biological signal itself.", width=13)

    # ------------------------------------------------------------ methods
    rp.h("3. Modelling approach (technical)")
    rp.bullets([
        "Strict separation of development and test data: every choice (features, preprocessing, model type, "
        f"hyperparameters, threshold) is made on the development set ({DEV}) only; the single-blind test set ({TEST}) is "
        "scored once at the end, as in the original study.",
        f"Nested cross-validation on {DEV}: an inner 3-fold grid search tunes each pipeline; an outer 5-fold CV, stratified "
        f"by label and sequencing batch and repeated {R['settings']['outer_repeats']} times, scores it.",
        "All preprocessing inside the pipelines (median imputation, standardisation), so nothing is learned from held-out folds.",
        "Candidate pipelines: L2 logistic regression; sparse L1 logistic regression; ANOVA top-k feature selection + "
        "logistic regression; PCA dimensionality reduction + logistic regression; a control-based outlier-count summary "
        "(share of regions far above or below the training controls, which condenses rare tumour signal spread over many "
        "regions); histogram gradient boosting; and, on region methylation only (too slow on the larger input), an RBF "
        "support vector machine on PCA components and a random forest.",
        "Two feature inputs: region methylation only, and methylation plus all read-level features (read-level "
        "methylation, methylated-fragment fraction in all and in short fragments, methylation haplotype load, discordant "
        "reads, entropy, short-fragment fraction, and sample-level fragment length).",
        "Pre-specified final models: the single best pipeline by nested AUC, and the mean of the three best distinct pipelines.",
        f"Metrics: ROC-AUC and sensitivity at 95% specificity, with the threshold fixed from {DEV} out-of-fold scores; "
        f"{R['settings']['boot']} bootstrap resamples of {TEST} for 95% CIs; AUC within each sequencing batch.",
        f"Controls: a covariates-only model scored on {TEST}; a {perm['n']}-shuffle label-permutation test and CV within "
        f"each batch, both on all samples ({DEV} and {TEST}) and not used for any choice.",
    ])

    # ------------------------------------------------------------ selection
    rp.h("4. Model selection on the development set")
    rp.p(f"The best pipeline was {s['best_candidate']} on {s['best_input'].lower()} (nested AUC "
         f"{num(s['best_nested_auc'], 3)}). The three best distinct pipelines, averaged as an ensemble, were: "
         f"{'; '.join(s['ensemble_members'])}.")
    rp.figure(res / "fig04_model_selection.png", f"Figure 4. Nested cross-validated AUC of every pipeline on {DEV} "
              "(mean ± SD over outer repeats).")
    top = nested.head(10).copy()
    top["nested AUC"] = top["nested_auc"].map(lambda v: num(v, 3)) + " ± " + top["nested_auc_sd"].map(lambda v: num(v, 3))
    top["sens @95% spec"] = top["sens_at_95spec"].map(pct)
    rp.table(top[["candidate", "input", "nested AUC", "sens @95% spec", "typical_params"]], fontsize=7.5)

    # ------------------------------------------------------------ held-out
    rp.h("5. Results on the independent test set")
    show = ho.sort_values("auc_test", ascending=False)
    pre = pd.concat([ho[ho["role"] == "selected"], ho[ho["role"] == "pre-specified ensemble"]])
    tab = pre.assign(AUC=lambda x: x["auc_test"].map(num) + " (" + x["auc_ci_low"].map(num) + "–" + x["auc_ci_high"].map(num) + ")",
                     sens=lambda x: x["sens_at_95spec"].map(pct) + " (" + x["sens_ci_low"].map(pct) + "–" + x["sens_ci_high"].map(pct) + ")",
                     spec=lambda x: x["spec_reached"].map(pct))
    for b in batches:
        tab[f"AUC {b}"] = tab[f"auc_test_{b}"].map(lambda v: num(v) if v == v else "–")
    rp.table(tab[["role", "candidate", "input", "AUC", "sens", "spec", *[f"AUC {b}" for b in batches]]]
             .rename(columns={"sens": "sens @95% spec", "spec": "specificity reached"}), fontsize=8)
    rp.figure(res / "fig05_heldout.png", "Figure 5. Left: ROC curves on the test set. Right: predicted scores of the "
              "selected model; the dashed line is the 95%-specificity threshold fixed on the development set.")
    rp.p(batch_line)
    batch_figure(res / "fig07_batch.png", sel, ens, batches, scorable, test_cancers, cov_only, cv_within)
    rp.figure(res / "fig07_batch.png", "Figure 6. Test-set AUC overall and within each sequencing batch (n.e.: not "
              "estimable, too few cancers), and cross-validated AUC inside each batch alone; dashed: covariates only.", width=15)
    rp.p("All other pipelines on the test set (exploratory; not used for selection):", bold=True)
    ex = show.assign(AUC=lambda x: x["auc_test"].map(num), sens=lambda x: x["sens_at_95spec"].map(pct))
    rp.table(ex[["role", "candidate", "input", "AUC", "sens"]].rename(columns={"sens": "sens @95% spec"}), fontsize=7.5)

    # ------------------------------------------------------------ published classifier
    if bench:
        rp.h("Comparison with the published classifier")
        rp.p(bench_line)
        rows = []
        for key, name in [("all", f"All {TEST}"), *[(b, f"Batch {b}") for b in batches if b in bench]]:
            v = bench[key]
            rows.append([name, f"{v['n']} ({v['cancers']} cancers)", pct(v["published_sens"]), pct(v["published_spec"]),
                         pct(v["our_sens_at_published_spec"]) if v["our_sens_at_published_spec"] is not None else "–",
                         num(v["our_auc"]) if v["our_auc"] is not None else "–"])
        rp.table(pd.DataFrame(rows, columns=["group", "samples", "published sensitivity", "published specificity",
                                             "our sensitivity at that specificity", "our AUC"]), fontsize=8)
        rp.p("The published calls are binary, so only one operating point can be compared. Our threshold here is set on "
             "the test controls to match the published specificity, which slightly favours our model.", italic=True,
             color=MUTED, size=9)

    # ------------------------------------------------------------ robustness
    rp.h("6. Robustness and controls")
    rp.p(f"Permutation test: the selected configuration reaches a cross-validated AUC of {num(perm['real_auc'])} on all "
         f"samples with the real labels; on {perm['n']} label shuffles it averages {num(perm['null_mean'])} (maximum {num(perm['null_max'])}), "
         f"p = {perm['p']:.3f}.")
    rp.figure(res / "fig06_permutation.png", "Figure 7. Permutation test of the selected model.", width=12)
    rob_show = rob.copy()
    for c in rob_show.columns:
        rob_show[c] = rob_show[c].map(lambda v: num(v) if isinstance(v, float) and v == v else ("" if v is None or v != v else v))
    keep = [c for c in ["auc_test", "sens_at_95spec", "cv_auc"] if c in rob_show.columns]
    rp.table(rob_show[keep].reset_index().rename(columns={"index": "analysis", "auc_test": "test AUC",
                                                          "sens_at_95spec": "sens @95% spec", "cv_auc": "CV AUC"}), fontsize=8)

    # ------------------------------------------------------------ regions
    rp.h("7. Which regions carry the signal?")
    rp.figure(res / "fig08_top_regions.png", "Figure 8. The 20 regions that separate the groups best on their own, with "
              "the methylation difference and the region AUC.", width=13)
    tr = pd.DataFrame(R["top_regions"]).head(15).rename(columns={"index": "region_id"})
    cols = [c for c in ["region_id", "chrom", "start0", "end", "auc", "delta", "fdr"] if c in tr.columns]
    tr = tr[cols].assign(auc=lambda x: x["auc"].map(lambda v: num(v, 3)), delta=lambda x: x["delta"].map(lambda v: f"{100 * v:+.1f} pp"),
                         fdr=lambda x: x["fdr"].map(lambda v: f"{v:.1e}"))
    rp.table(tr, fontsize=8)

    # ------------------------------------------------------------ limitations / next
    rp.h("8. Limitations")
    rp.bullets([
        "One study, one laboratory protocol: performance on data from other sites is unknown.",
        comp_line,
        f"Sequencing batches differ: the development set has {int(dev_cancers.get(small, 0))} cancer(s) from the small "
        f"batch {small}, which also holds {benign.get(small, 0)} of the {sum(benign.values())} benign-nodule controls"
        + (", and the model does not transfer to that batch." if small in fails else "."),
        "The test set is moderate in size, so the confidence intervals are wide, especially for sensitivity at 95% specificity.",
        "Hyperparameter grids are small; larger searches may help.",
        "Fragment features depend on library preparation and should be interpreted with care.",
    ])
    rp.h("9. Recommended next steps")
    rp.bullets([
        "Train with cancers from both sequencing batches (e.g. a pooled, batch-stratified cross-validation of all samples, "
        f"or domain adaptation), since the {DEV} set has only {int(dev_cancers.get(small, 0))} cancer(s) from batch {small}.",
        *(["Close the gap to the published classifier at high specificity: tumour-specific fragment features and "
           "operating-point-aware training (e.g. partial AUC at high specificity)."]
          if bench and (bench["all"]["our_sens_at_published_spec"] or 0) < bench["all"]["published_sens"] else []),
        "Test on an external cohort (another study or our own samples).",
        "Annotate the top regions with genes and pathways; explain individual predictions (e.g. SHAP).",
        "Explore richer read-level models (fragment-level deconvolution of tumour fraction).",
    ])

    # ------------------------------------------------------------ appendix
    rp.h("Appendix A. Glossary for non-technical readers")
    rp.table(pd.DataFrame([
        ["AUC", "Probability that the model scores a random cancer higher than a random control. 0.5 = chance, 1.0 = perfect."],
        ["Sensitivity", "Share of people with cancer that the test flags."],
        ["Specificity", "Share of people without cancer that the test correctly clears. 95% means 5 false alarms per 100."],
        ["Cross-validation", "Repeatedly train on part of the data and test on the rest, to estimate performance honestly."],
        ["Held-out test set", "Data kept aside and used only once, at the end, to measure the final model."],
        ["Confidence interval", "Range in which the true value is likely to lie, given the sample size."],
        ["Batch", "A group of samples sequenced together; batches can differ for technical reasons."],
        ["Methylation region", "A stretch of DNA whose chemical tags (methylation) are measured; 2,473 regions in the panel."],
    ], columns=["Term", "Meaning"]), widths=[4, 12.5])
    rp.h("Appendix B. Confounders (each variable alone)")
    cf = conf.reset_index().rename(columns={"index": "variable"})
    rp.table(cf.assign(abs_auc=cf["abs_auc"].map(num), p=cf["p"].map(lambda v: f"{v:.2g}"),
                       control_median=cf["control_median"].map(lambda v: f"{v:.3g}"),
                       cancer_median=cf["cancer_median"].map(lambda v: f"{v:.3g}"))[
        ["variable", "abs_auc", "p", "control_median", "cancer_median"]], fontsize=8)
    rp.p(f"Analysis runtime: {R['runtime_min']:.0f} minutes. Code: scripts/08_ml_analysis.py, scripts/09_published_benchmark.py and "
         "scripts/10_report.py.",
         italic=True, color=MUTED, size=9)

    out = res / "ML_Report.docx"
    rp.doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
